# -*- coding: utf-8 -*-
"""Конвертация docs/report.md в DOCX (стили: Times New Roman 14, GOST-подобные заголовки)."""
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = r"D:\大四\de\LabRabota1\docs\report.md"
DST = r"D:\大四\de\LabRabota1\docs\report.docx"

doc = Document()

# Базовый стиль
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(14)
style.paragraph_format.line_spacing = 1.5
style.paragraph_format.first_line_indent = Cm(1.25)

# Поля страницы
for section in doc.sections:
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)


def add_runs(par, text):
    """Разбор инлайн-разметки **bold** и `code`."""
    tokens = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = par.add_run(tok[2:-2])
            r.bold = True
        elif tok.startswith("`") and tok.endswith("`"):
            r = par.add_run(tok[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(12)
            r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        else:
            par.add_run(tok)


def flush_table(rows):
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(ncol):
            cell = t.cell(i, j)
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.first_line_indent = Cm(0)
            txt = row[j] if j < len(row) else ""
            txt = txt.replace("**", "")
            run = par.add_run(txt)
            run.font.size = Pt(12)
            if i == 0:
                run.bold = True
    doc.add_paragraph()


with open(SRC, encoding="utf-8") as f:
    lines = f.read().splitlines()

i = 0
in_code = False
code_buf = []
table_buf = []

while i < len(lines):
    line = lines[i]

    # Блок кода
    if line.strip().startswith("```"):
        if in_code:
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(6)
            for cl in code_buf:
                r = p.add_run(cl + "\n")
                r.font.name = "Consolas"
                r.font.size = Pt(10)
            shade = p._p.get_or_add_pPr()
            from docx.oxml.ns import qn
            from docx.oxml import OxmlElement
            sh = OxmlElement("w:shd")
            sh.set(qn("w:val"), "clear")
            sh.set(qn("w:fill"), "F2F2F2")
            shade.append(sh)
            code_buf = []
            in_code = False
        else:
            in_code = True
        i += 1
        continue
    if in_code:
        code_buf.append(line)
        i += 1
        continue

    # Таблица markdown
    if line.strip().startswith("|"):
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
            i += 1
            continue  # строка-разделитель
        table_buf.append(cells)
        i += 1
        continue
    else:
        if table_buf:
            flush_table(table_buf)
            table_buf = []

    s = line.strip()

    if not s:
        i += 1
        continue

    if s == "---":
        i += 1
        continue

    # HTML-комментарии — пропускаем, но оставляем placeholder-текст
    if s.startswith("<!--") and s.endswith("-->"):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        r = p.add_run(f"[{s[4:-3].strip()}]")
        r.italic = True
        r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
        i += 1
        continue

    m = re.match(r"^(#{1,4})\s+(.*)", s)
    if m:
        level = len(m.group(1))
        text = m.group(2).replace("**", "")
        h = doc.add_heading(text, level=min(level, 4))
        h.paragraph_format.first_line_indent = Cm(0)
        for r in h.runs:
            r.font.name = "Times New Roman"
            r.font.color.rgb = RGBColor(0, 0, 0)
            r.font.size = Pt({1: 18, 2: 16, 3: 14, 4: 14}[min(level, 4)])
            r.bold = True
        if level == 1:
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        i += 1
        continue

    if s.startswith(">"):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.left_indent = Cm(1.0)
        add_runs(p, s.lstrip("> ").strip())
        for r in p.runs:
            r.italic = True
            r.font.size = Pt(12)
        i += 1
        continue

    mli = re.match(r"^[-*]\s+(.*)", s)
    if mli:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.first_line_indent = Cm(0)
        add_runs(p, mli.group(1))
        i += 1
        continue

    mnum = re.match(r"^\d+[.)]\s+(.*)", s)
    if mnum:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.first_line_indent = Cm(0)
        add_runs(p, mnum.group(1))
        i += 1
        continue

    # Обычный абзац
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_runs(p, s)
    i += 1

if table_buf:
    flush_table(table_buf)

doc.save(DST)
print("saved:", DST)
